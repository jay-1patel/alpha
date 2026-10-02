import os
import sys
import json
import re
import logging
import shutil
from pathlib import Path
from typing import Optional, List, Dict, Any

import fitz  # PyMuPDF
import pdfplumber
import pytesseract
from docx import Document
import openpyxl
from PIL import Image

from database import init_db, save_chunks

# ----------------------------------------------------------------------
# Logging Setup
# ----------------------------------------------------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)
logger = logging.getLogger("extract")

# ----------------------------------------------------------------------
# Configuration (env-var driven, with sane fallbacks)
# ----------------------------------------------------------------------

# Tesseract path — configurable via env var so this works on any machine,
# not just yours. Falls back to auto-detection via shutil.which(),
# then to a Windows-typical default path if nothing else is found.
TESSERACT_CMD = os.environ.get("TESSERACT_CMD") or shutil.which("tesseract")

if not TESSERACT_CMD:
    _default_windows_path = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    if os.path.exists(_default_windows_path):
        TESSERACT_CMD = _default_windows_path

if TESSERACT_CMD:
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD
    logger.info(f"Using Tesseract at: {TESSERACT_CMD}")
else:
    logger.warning(
        "Tesseract executable not found. OCR on scanned/image-only PDF pages "
        "will fail. Set the TESSERACT_CMD environment variable to the full "
        "path of tesseract.exe."
    )

CHUNK_SIZE = int(os.environ.get("CHUNK_SIZE", 500))
CHUNK_OVERLAP = int(os.environ.get("CHUNK_OVERLAP", 60))

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.environ.get("UPLOAD_DIR", os.path.join(BASE_DIR, "uploaded_files"))
EXTRACTED_DIR = os.environ.get("EXTRACTED_DIR", os.path.join(UPLOAD_DIR, "extracted"))

os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(EXTRACTED_DIR, exist_ok=True)

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".doc", ".txt", ".xlsx", ".xls")


# ----------------------------------------------------------------------
# Text Splitting Helpers
# ----------------------------------------------------------------------

def split_qa(text: str) -> List[str]:
    """Split text into chunks at Q:/Question:/Ques: boundaries."""
    if not text:
        return []
    parts = re.split(
        r'(?=(?:\bQ\b|\bQuestion\b|\bquestion\b|\bQues\b|\bques\b):\s)', text
    )
    return [p.strip() for p in parts if p.strip()]


def split_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    """Fallback fixed-size chunker with overlap for text that has no Q/A markers."""
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0
    text_len = len(text)

    while start < text_len:
        end = start + chunk_size
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        next_start = end - overlap
        # Guard against infinite loop if overlap >= chunk_size
        start = next_start if next_start > start else end

    return chunks


def chunk_paragraph(text: str) -> List[str]:
    """Common helper: try Q/A splitting first, fall back to fixed-size chunking."""
    splits = split_qa(text)
    if len(splits) <= 1 and len(text) > CHUNK_SIZE:
        return split_text(text)
    return splits


# ----------------------------------------------------------------------
# OCR
# ----------------------------------------------------------------------

def ocr_page(page) -> str:
    """Render a PDF page to an image and run OCR on it."""
    if not TESSERACT_CMD:
        logger.warning("Skipping OCR — tesseract executable not configured.")
        return ""
    try:
        pix = page.get_pixmap(dpi=300)
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        text = pytesseract.image_to_string(img)
        return text.strip()
    except Exception as e:
        logger.error(f"OCR failed on page: {e}")
        return ""


# ----------------------------------------------------------------------
# PDF Extraction
# ----------------------------------------------------------------------

def extract_tables_pdf(filepath: str) -> List[Dict[str, Any]]:
    tables = []
    try:
        with pdfplumber.open(filepath) as pdf:
            for i, page in enumerate(pdf.pages):
                try:
                    page_tables = page.extract_tables()
                except Exception as e:
                    logger.warning(f"Table extraction failed on page {i + 1}: {e}")
                    continue

                for t_idx, table in enumerate(page_tables):
                    if table:
                        cleaned = [[cell if cell else "" for cell in row] for row in table]
                        tables.append({
                            "page": i + 1,
                            "table_index": t_idx + 1,
                            "data": cleaned,
                        })
    except Exception as e:
        logger.error(f"Failed to open PDF for table extraction ({filepath}): {e}")

    return tables


def extract_pdf(filepath: str) -> Dict[str, Any]:
    results = {"file": os.path.basename(filepath), "type": "pdf", "pages": [], "tables": []}

    tables = extract_tables_pdf(filepath)
    results["tables"] = tables
    table_pages = {t["page"] for t in tables}

    try:
        doc = fitz.open(filepath)
    except Exception as e:
        logger.error(f"Failed to open PDF ({filepath}): {e}")
        return results

    try:
        for i, page in enumerate(doc):
            page_num = i + 1
            try:
                text = page.get_text("text").strip()
            except Exception as e:
                logger.warning(f"Text extraction failed on page {page_num}: {e}")
                text = ""

            if not text and page_num not in table_pages:
                ocr_text = ocr_page(page)
                results["pages"].append({"page": page_num, "text": ocr_text, "source": "ocr"})
            elif text:
                results["pages"].append({"page": page_num, "text": text, "source": "text"})
            else:
                results["pages"].append({"page": page_num, "text": "", "source": "empty"})
    finally:
        doc.close()

    return results


# ----------------------------------------------------------------------
# Word Extraction
# ----------------------------------------------------------------------

def extract_word(filepath: str) -> Dict[str, Any]:
    results = {"file": os.path.basename(filepath), "type": "docx", "paragraphs": [], "tables": []}

    try:
        doc = Document(filepath)
    except Exception as e:
        logger.error(f"Failed to open DOCX ({filepath}): {e}")
        return results

    for para in doc.paragraphs:
        if para.text.strip():
            results["paragraphs"].append(para.text.strip())

    for t_idx, table in enumerate(doc.tables):
        try:
            table_data = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            results["tables"].append({"table_index": t_idx + 1, "data": table_data})
        except Exception as e:
            logger.warning(f"Failed to extract table {t_idx + 1}: {e}")

    return results


# ----------------------------------------------------------------------
# TXT Extraction
# ----------------------------------------------------------------------

def extract_txt(filepath: str) -> Dict[str, Any]:
    results = {"file": os.path.basename(filepath), "type": "txt", "paragraphs": []}

    try:
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            for line in f:
                stripped = line.strip()
                if stripped:
                    results["paragraphs"].append(stripped)
    except Exception as e:
        logger.error(f"Failed to read TXT ({filepath}): {e}")

    return results


# ----------------------------------------------------------------------
# Excel Extraction
# ----------------------------------------------------------------------

def extract_excel(filepath: str) -> Dict[str, Any]:
    results = {"file": os.path.basename(filepath), "type": "xlsx", "sheets": [], "tables": []}

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
    except Exception as e:
        logger.error(f"Failed to open Excel file ({filepath}): {e}")
        return results

    try:
        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            sheet_data = []

            for row in ws.iter_rows(values_only=True):
                row_vals = [str(cell) if cell is not None else "" for cell in row]
                if any(v.strip() for v in row_vals):
                    sheet_data.append(row_vals)

            if sheet_data:
                results["sheets"].append({"sheet_name": sheet_name, "data": sheet_data})
                results["tables"].append({
                    "table_index": sheet_idx + 1,
                    "data": sheet_data,
                    "sheet_name": sheet_name,
                })
    finally:
        wb.close()

    return results


# ----------------------------------------------------------------------
# Dispatcher
# ----------------------------------------------------------------------

def process_file(filepath: str) -> Optional[Dict[str, Any]]:
    """Route a file to the appropriate extractor based on its extension."""
    ext = Path(filepath).suffix.lower()

    if not os.path.exists(filepath):
        logger.error(f"File not found: {filepath}")
        return None

    try:
        if ext == ".pdf":
            return extract_pdf(filepath)
        elif ext in (".docx", ".doc"):
            return extract_word(filepath)
        elif ext == ".txt":
            return extract_txt(filepath)
        elif ext in (".xlsx", ".xls"):
            return extract_excel(filepath)
        else:
            logger.warning(f"Unsupported file type: {ext}")
            return None
    except Exception as e:
        logger.exception(f"Unexpected error processing file {filepath}: {e}")
        return None


# ----------------------------------------------------------------------
# Convert Extraction Results -> DB Chunks
# ----------------------------------------------------------------------

def results_to_chunks(results: Dict[str, Any]) -> List[Dict[str, Any]]:
    chunks = []
    filename = results["file"]
    file_type = results["type"]

    if file_type == "pdf":
        for page in results["pages"]:
            if page["text"]:
                content_type = "ocr" if page["source"] == "ocr" else "paragraph"
                for s in chunk_paragraph(page["text"]):
                    chunks.append({
                        "source_file": filename,
                        "content_type": content_type,
                        "content": s,
                        "page_number": page["page"],
                    })

        for table in results["tables"]:
            table_text = "\n".join(" | ".join(row) for row in table["data"])
            chunks.append({
                "source_file": filename,
                "content_type": "table",
                "content": table_text,
                "page_number": table["page"],
            })

    elif file_type == "docx":
        full_text = "\n".join(results["paragraphs"])
        if full_text.strip():
            for s in chunk_paragraph(full_text):
                chunks.append({
                    "source_file": filename,
                    "content_type": "paragraph",
                    "content": s,
                    "page_number": 1,
                })

        for table in results["tables"]:
            table_text = "\n".join(" | ".join(row) for row in table["data"])
            chunks.append({
                "source_file": filename,
                "content_type": "table",
                "content": table_text,
                "page_number": 1,
            })

    elif file_type == "txt":
        full_text = "\n".join(results["paragraphs"])
        if full_text.strip():
            for s in chunk_paragraph(full_text):
                chunks.append({
                    "source_file": filename,
                    "content_type": "paragraph",
                    "content": s,
                    "page_number": 1,
                })

    elif file_type == "xlsx":
        for sheet in results["sheets"]:
            for row in sheet["data"]:
                row_text = " | ".join(row)
                if row_text.strip():
                    chunks.append({
                        "source_file": filename,
                        "content_type": "table",
                        "content": row_text,
                        "page_number": 1,
                    })

    return chunks


# ----------------------------------------------------------------------
# Human-readable .txt Report
# ----------------------------------------------------------------------

def format_output(results: Dict[str, Any]) -> str:
    lines = []
    lines.append("=" * 60)
    lines.append(f"FILE: {results['file']}")
    lines.append(f"TYPE: {results['type']}")
    lines.append("=" * 60)

    file_type = results["type"]

    if file_type == "pdf":
        for page in results["pages"]:
            lines.append(f"\n--- Page {page['page']} ({page['source']}) ---\n")
            lines.append(page["text"] if page["text"] else "[No content]")

        if results["tables"]:
            lines.append(f"\n{'=' * 60}")
            lines.append("TABLES")
            lines.append("=" * 60)
            for table in results["tables"]:
                lines.append(f"\n[Table {table['table_index']} from Page {table['page']}]")
                for row in table["data"]:
                    lines.append(" | ".join(row))

    elif file_type == "docx":
        lines.append("\n--- Paragraphs ---\n")
        for para in results["paragraphs"]:
            lines.append(para)
            lines.append("")

        if results["tables"]:
            lines.append(f"\n{'=' * 60}")
            lines.append("TABLES")
            lines.append("=" * 60)
            for table in results["tables"]:
                lines.append(f"\n[Table {table['table_index']}]")
                for row in table["data"]:
                    lines.append(" | ".join(row))

    elif file_type == "txt":
        lines.append("\n--- Content ---\n")
        for para in results["paragraphs"]:
            lines.append(para)
            lines.append("")

    elif file_type == "xlsx":
        for sheet in results["sheets"]:
            lines.append(f"\n--- Sheet: {sheet['sheet_name']} ---\n")
            for row in sheet["data"]:
                lines.append(" | ".join(row))
            lines.append("")

    return "\n".join(lines)


# ----------------------------------------------------------------------
# Single-file pipeline — designed to be called from FastAPI via
# run_in_threadpool(process_uploaded_file, filepath) since every step
# here is blocking / CPU-bound (OCR, PDF parsing, Excel parsing, etc.)
# ----------------------------------------------------------------------

def process_uploaded_file(filepath: str, save_to_db: bool = True) -> Dict[str, Any]:
    """
    Full pipeline for a single uploaded file:
      1. Extract content (text/tables/OCR as needed)
      2. Write human-readable .txt and .json reports to EXTRACTED_DIR
      3. Convert to chunks
      4. Optionally persist chunks to the database

    Returns a summary dict — safe to serialize directly as a FastAPI response.
    Raises no exceptions for expected failure modes (bad/corrupt file);
    instead returns a dict with "success": False and an "error" message.
    """
    filename = os.path.basename(filepath)
    ext = Path(filepath).suffix.lower()

    if ext not in SUPPORTED_EXTENSIONS:
        return {
            "success": False,
            "file": filename,
            "error": f"Unsupported file type: {ext}",
        }

    logger.info(f"Processing uploaded file: {filename}")

    try:
        results = process_file(filepath)
    except Exception as e:
        logger.exception(f"Extraction crashed for {filename}")
        return {"success": False, "file": filename, "error": str(e)}

    if not results:
        return {
            "success": False,
            "file": filename,
            "error": "Extraction returned no results (see server logs for details).",
        }

    stem = Path(filename).stem

    # Write human-readable report
    txt_report_path = os.path.join(EXTRACTED_DIR, f"{stem}_extracted.txt")
    try:
        with open(txt_report_path, "w", encoding="utf-8") as f:
            f.write(format_output(results))
    except Exception as e:
        logger.warning(f"Could not write text report for {filename}: {e}")
        txt_report_path = None

    # Write JSON dump
    json_report_path = os.path.join(EXTRACTED_DIR, f"{stem}_extracted.json")
    try:
        with open(json_report_path, "w", encoding="utf-8") as f:
            json.dump(results, f, indent=2, ensure_ascii=False)
    except Exception as e:
        logger.warning(f"Could not write JSON report for {filename}: {e}")
        json_report_path = None

    # Build chunks
    try:
        chunks = results_to_chunks(results)
    except Exception as e:
        logger.exception(f"Chunking failed for {filename}")
        return {"success": False, "file": filename, "error": f"Chunking failed: {e}"}

    # Persist to DB
    saved = False
    if save_to_db and chunks:
        try:
            init_db()
            save_chunks(chunks)
            saved = True
        except Exception as e:
            logger.exception(f"Failed to save chunks to DB for {filename}")
            return {
                "success": False,
                "file": filename,
                "error": f"Chunks extracted but DB save failed: {e}",
                "chunk_count": len(chunks),
            }

    logger.info(f"Finished processing {filename}: {len(chunks)} chunks extracted")

    return {
        "success": True,
        "file": filename,
        "file_type": results["type"],
        "chunk_count": len(chunks),
        "saved_to_db": saved,
        "text_report": txt_report_path,
        "json_report": json_report_path,
    }


# ----------------------------------------------------------------------
# CLI Entry Point (batch-process everything in UPLOAD_DIR)
# ----------------------------------------------------------------------

def main():
    files = [
        f for f in os.listdir(UPLOAD_DIR)
        if Path(f).suffix.lower() in SUPPORTED_EXTENSIONS
    ]

    if not files:
        logger.info(f"No supported files found in {UPLOAD_DIR}")
        return

    init_db()
    logger.info(f"Found {len(files)} file(s) to process")

    all_chunks = []
    summary = []

    for filename in files:
        filepath = os.path.join(UPLOAD_DIR, filename)
        logger.info(f"Processing: {filename}")

        result = process_uploaded_file(filepath, save_to_db=False)
        summary.append(result)

        if result.get("success"):
            # Re-run chunk extraction here just to accumulate for a single
            # batched save_chunks() call at the end (avoids N separate DB writes)
            results = process_file(filepath)
            if results:
                chunks = results_to_chunks(results)
                all_chunks.extend(chunks)
                logger.info(f"  -> {len(chunks)} chunks extracted")
        else:
            logger.error(f"  -> Failed: {result.get('error')}")

    if all_chunks:
        save_chunks(all_chunks)
        logger.info(f"Saved {len(all_chunks)} chunks to faq_dataset table")

    logger.info(f"Done. Extracted files saved in: {EXTRACTED_DIR}")

    failed = [s for s in summary if not s.get("success")]
    if failed:
        logger.warning(f"{len(failed)} file(s) failed to process:")
        for f in failed:
            logger.warning(f"  - {f['file']}: {f.get('error')}")


if __name__ == "__main__":
    main()